"""Document parsers: extract per-page/per-slide text from PDF and PPTX.

- PDFs are parsed with PyMuPDF (``fitz``).
- PPTX files are parsed with python-pptx, including speaker notes.
Every page/slide is returned with its 1-based number so chunk metadata stays
useful for citation on the results page.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class PageContent:
    """Text extracted from a single page (PDF) or slide (PPTX)."""

    text: str
    page_number: int


SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".pptx": "pptx"}


def parse_document(path: str | Path, source_type: str) -> list[PageContent]:
    """Dispatch to the right parser based on ``source_type`` ("pdf"|"pptx")."""
    resolved_type = source_type.lower().lstrip(".")
    if resolved_type == "pdf":
        return parse_pdf(path)
    if resolved_type == "pptx":
        return parse_pptx(path)
    raise ValueError(f"Unsupported source type '{source_type}'. Expected 'pdf' or 'pptx'.")


def parse_pdf(path: str | Path) -> list[PageContent]:
    """Extract text from every page of a PDF using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is not installed. Run: pip install -r backend/requirements.txt") from exc

    pages: list[PageContent] = []
    document = fitz.open(str(path))
    try:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(PageContent(text=text, page_number=page_number))
    finally:
        document.close()
    if not pages:
        raise ValueError("No extractable text found in this PDF (it may be a scanned image PDF).")
    return pages


def parse_pptx(path: str | Path) -> list[PageContent]:
    """Extract text and speaker notes from every slide of a PPTX file."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is not installed. Run: pip install -r backend/requirements.txt") from exc

    presentation = Presentation(str(path))
    pages: list[PageContent] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(_text_frame_content(shape.text_frame))
            if getattr(shape, "has_table", False):
                parts.append(_table_content(shape.table))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        text = "\n".join(p for p in parts if p).strip()
        if text:
            pages.append(PageContent(text=text, page_number=slide_number))
    if not pages:
        raise ValueError("No text found in this PPTX (empty slides).")
    return pages


def _text_frame_content(text_frame) -> str:
    lines: list[str] = []
    for paragraph in text_frame.paragraphs:
        runs = [run.text for run in paragraph.runs if run.text]
        if runs:
            lines.append("".join(runs).strip())
    return "\n".join(lines).strip()


def _table_content(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            value = cell.text.strip()
            if value:
                cells.append(value)
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows).strip()